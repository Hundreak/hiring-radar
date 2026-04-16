# src/hiring_radar/services/cv_parser_engine.py
from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import logging
import mimetypes
import os
import re
import tempfile
import time
import unicodedata
from collections import Counter
from collections.abc import Iterable, Sequence
from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

try:
    import fitz  # type: ignore[import-not-found]
except Exception:  # pragma: no cover - optional dependency
    fitz = None

try:
    import pdfplumber  # type: ignore[import-not-found]
except Exception:  # pragma: no cover - optional dependency
    pdfplumber = None

try:
    import pytesseract  # type: ignore[import-not-found]
except Exception:  # pragma: no cover - optional dependency
    pytesseract = None

try:
    import easyocr  # type: ignore[import-not-found]
except Exception:  # pragma: no cover - optional dependency
    easyocr = None

try:
    from PIL import Image, ImageFilter, ImageOps
except Exception:  # pragma: no cover - optional dependency
    Image = None
    ImageFilter = None
    ImageOps = None

try:
    from docx import Document
except Exception:  # pragma: no cover - optional dependency
    Document = None

LOGGER = logging.getLogger("hiring_radar.cv_parser_engine")

PARSER_VERSION = "universal-cv-engine-v1"
DEFAULT_OCR_LANGUAGES = ("eng",)
CONTACT_SCAN_LINE_LIMIT = 20
SUMMARY_SCAN_LINE_LIMIT = 12
MAX_SUMMARY_CHARS = 1_500
MIN_SPARSE_TEXT_LENGTH = 80
SECTION_NAME_MAX_WORDS = 6
MAX_NAME_TOKENS = 6
MAX_HEADLINE_LENGTH = 120
MAX_BLOCKS_PER_SECTION = 100
MAX_BULLETS_PER_ENTRY = 12
MAX_SKILLS = 80
MAX_TARGET_ROLES = 20
MAX_URLS = 8
MAX_CERTIFICATIONS = 25

EMAIL_RE = re.compile(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s()./-]{7,}\d)")
URL_RE = re.compile(r"https?://\S+|www\.\S+", re.IGNORECASE)
LINKEDIN_RE = re.compile(r"linkedin\.com", re.IGNORECASE)
GITHUB_RE = re.compile(r"github\.com", re.IGNORECASE)
DATE_RANGE_RE = re.compile(
    r"(?P<start>(?:\b(?:19|20)\d{2}\b|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)"
    r"[a-z]*\s+(?:19|20)\d{2}|(?:0?[1-9]|1[0-2])[./-](?:19|20)?\d{2}))"
    r"\s*(?:-|–|—|to|until|/|\u2013|\u2014)\s*"
    r"(?P<end>(?:\b(?:19|20)\d{2}\b|present|current|ongoing|today|heute|devam|"
    r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+(?:19|20)\d{2}|"
    r"(?:0?[1-9]|1[0-2])[./-](?:19|20)?\d{2}))",
    re.IGNORECASE,
)
YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
BULLET_PREFIX_RE = re.compile(r"^(?:[-•*▪◦‣⁃]+|\d+[.)])\s+")
PIPE_SPLIT_RE = re.compile(r"\s+[|·•‧▪/]+\s+")
COMMA_SPLIT_RE = re.compile(r"\s*,\s*")
SEMICOLON_SPLIT_RE = re.compile(r"\s*;\s*")
WHITESPACE_RE = re.compile(r"[ \t]+")
MULTIBREAK_RE = re.compile(r"\n{3,}")
NON_ALNUM_RE = re.compile(r"[^\w\s@+./#-]+", re.UNICODE)

REMOTE_KEYWORDS = frozenset(
    {
        "remote",
        "hybrid",
        "onsite",
        "on-site",
        "uzaktan",
        "hibrit",
        "ofiste",
        "remote first",
        "vor ort",
        "hybrid remote",
    }
)

ROLE_HINT_KEYWORDS = frozenset(
    {
        "engineer",
        "developer",
        "analyst",
        "manager",
        "scientist",
        "architect",
        "consultant",
        "specialist",
        "lead",
        "administrator",
        "designer",
        "intern",
        "technician",
        "researcher",
        "product",
        "project",
        "marketing",
        "sales",
        "finance",
        "accountant",
        "mühendis",
        "geliştirici",
        "analist",
        "uzman",
        "yönetici",
        "tasarımcı",
        "ingenieur",
        "entwickler",
        "berater",
        "leiter",
        "spezialist",
        "forscher",
    }
)

DEGREE_HINTS = frozenset(
    {
        "bsc",
        "msc",
        "bs",
        "ms",
        "ba",
        "ma",
        "m.eng",
        "b.eng",
        "phd",
        "mba",
        "bachelor",
        "master",
        "doctorate",
        "associate",
        "lisans",
        "yüksek lisans",
        "universite",
        "üniversite",
        "university",
        "college",
        "institute",
        "academy",
        "diploma",
    }
)

KNOWN_LANGUAGE_NAMES = frozenset(
    {
        "english",
        "turkish",
        "türkçe",
        "turkce",
        "german",
        "deutsch",
        "french",
        "français",
        "francais",
        "spanish",
        "español",
        "espanol",
        "italian",
        "italiano",
        "portuguese",
        "português",
        "portugues",
        "arabic",
        "russian",
        "dutch",
        "chinese",
        "mandarin",
        "japanese",
        "korean",
        "polish",
        "ukrainian",
    }
)

LANGUAGE_PROFICIENCY_HINTS = frozenset(
    {
        "native",
        "fluent",
        "professional",
        "conversational",
        "intermediate",
        "advanced",
        "beginner",
        "basic",
        "elementary",
        "mother tongue",
        "ana dil",
        "akıcı",
        "profesyonel",
        "orta",
        "ileri",
        "başlangıç",
        "muttersprache",
        "fließend",
        "grundkenntnisse",
        "fortgeschritten",
        "c2",
        "c1",
        "b2",
        "b1",
        "a2",
        "a1",
    }
)

SKILL_LEXICON = frozenset(
    {
        "python",
        "java",
        "javascript",
        "typescript",
        "go",
        "golang",
        "rust",
        "c",
        "c++",
        "c#",
        ".net",
        "kotlin",
        "swift",
        "php",
        "ruby",
        "scala",
        "sql",
        "postgresql",
        "mysql",
        "sqlite",
        "mongodb",
        "redis",
        "elasticsearch",
        "fastapi",
        "django",
        "flask",
        "spring",
        "node.js",
        "react",
        "next.js",
        "vue",
        "angular",
        "html",
        "css",
        "tailwind",
        "docker",
        "kubernetes",
        "terraform",
        "ansible",
        "aws",
        "azure",
        "gcp",
        "linux",
        "git",
        "github actions",
        "gitlab ci",
        "jenkins",
        "pandas",
        "numpy",
        "scikit-learn",
        "tensorflow",
        "pytorch",
        "opencv",
        "nlp",
        "machine learning",
        "deep learning",
        "llm",
        "rag",
        "langchain",
        "airflow",
        "spark",
        "hadoop",
        "rabbitmq",
        "kafka",
        "grpc",
        "rest",
        "microservices",
        "embedded",
        "matlab",
        "simulink",
        "pcb",
        "altium",
        "verilog",
        "vhdl",
        "fpga",
        "cad",
        "sap",
        "salesforce",
        "figma",
        "photoshop",
        "excel",
        "power bi",
        "tableau",
        "jira",
        "notion",
        "scrum",
        "agile",
        "kanban",
        "selenium",
        "pytest",
        "unit testing",
        "integration testing",
        "ci/cd",
    }
)

SECTION_ALIASES: dict[str, frozenset[str]] = {
    "summary": frozenset(
        {
            "summary",
            "professional summary",
            "executive summary",
            "profile",
            "about",
            "objective",
            "career summary",
            "personal profile",
            "professional profile",
            "career profile",
            "özet",
            "profil",
            "profil özeti",
            "kariyer özeti",
            "zusammenfassung",
            "profilzusammenfassung",
            "kurzprofil",
        }
    ),
    "skills": frozenset(
        {
            "skills",
            "technical skills",
            "core skills",
            "tech stack",
            "technologies",
            "competencies",
            "expertise",
            "key skills",
            "core competencies",
            "technical competencies",
            "beceriler",
            "teknik beceriler",
            "yetkinlikler",
            "uzmanlık alanları",
            "fähigkeiten",
            "kenntnisse",
            "technologien",
            "kompetenzen",
            "schlüsselkompetenzen",
        }
    ),
    "experience": frozenset(
        {
            "experience",
            "work experience",
            "employment history",
            "professional experience",
            "career history",
            "work history",
            "employment",
            "iş deneyimi",
            "deneyim",
            "tecrübe",
            "profesyonel deneyim",
            "berufserfahrung",
            "arbeitserfahrung",
            "beruflicher werdegang",
        }
    ),
    "education": frozenset(
        {
            "education",
            "academic background",
            "academic history",
            "education history",
            "qualifications",
            "eğitim",
            "öğrenim",
            "akademik geçmiş",
            "ausbildung",
            "bildung",
            "akademischer hintergrund",
            "qualifikationen",
        }
    ),
    "languages": frozenset(
        {
            "languages",
            "language",
            "language skills",
            "language proficiency",
            "diller",
            "dil",
            "sprachen",
            "sprachkenntnisse",
        }
    ),
    "certifications": frozenset(
        {
            "certifications",
            "certification",
            "licenses",
            "licences",
            "certificates",
            "sertifikalar",
            "sertifikalar ve lisanslar",
            "zertifikate",
            "lizenzen",
        }
    ),
    "target_roles": frozenset(
        {
            "target roles",
            "target role",
            "desired position",
            "desired roles",
            "career objective",
            "target position",
            "hedef rol",
            "hedef roller",
            "hedef pozisyon",
            "hedef pozisyonlar",
            "zielposition",
            "zielrolle",
            "gewünschte position",
        }
    ),
    "preferences": frozenset(
        {
            "preferences",
            "work preferences",
            "location preferences",
            "remote preference",
            "work authorization",
            "tercihler",
            "çalışma tercihleri",
            "lokasyon tercihleri",
            "arbeitspräferenzen",
            "standortpräferenzen",
        }
    ),
}


def _normalize_alias_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    without_marks = "".join(
        character for character in normalized if not unicodedata.combining(character)
    )
    collapsed = WHITESPACE_RE.sub(" ", without_marks).strip()
    collapsed = NON_ALNUM_RE.sub(" ", collapsed)
    return WHITESPACE_RE.sub(" ", collapsed).strip().casefold()


NORMALIZED_SECTION_ALIASES: dict[str, frozenset[str]] = {
    canonical_name: frozenset(_normalize_alias_key(alias) for alias in aliases)
    for canonical_name, aliases in SECTION_ALIASES.items()
}


DocumentFileType = Literal["pdf", "docx", "png", "jpg", "jpeg", "image", "unknown"]
ParseStatus = Literal["parsed", "partial", "empty", "failed"]
QualityLabel = Literal["high", "medium", "low"]
WarningSeverity = Literal["info", "warning", "error"]


class ParserConfig(BaseModel):
    """Runtime configuration for the universal CV parser engine."""

    model_config = ConfigDict(extra="forbid")

    ocr_languages: tuple[str, ...] = DEFAULT_OCR_LANGUAGES
    preferred_ocr_backends: tuple[str, ...] = ("pytesseract", "easyocr")
    pdf_sparse_text_threshold: int = MIN_SPARSE_TEXT_LENGTH
    max_summary_chars: int = MAX_SUMMARY_CHARS
    async_concurrency: int = 4
    process_workers: int = 4
    render_scale: float = 2.0
    include_raw_text: bool = True
    include_normalized_text: bool = True
    max_urls: int = MAX_URLS
    max_skills: int = MAX_SKILLS
    max_target_roles: int = MAX_TARGET_ROLES
    max_certifications: int = MAX_CERTIFICATIONS


class ParseWarning(BaseModel):
    """Structured warning emitted during extraction or parsing."""

    model_config = ConfigDict(extra="forbid")

    code: str
    message: str
    severity: WarningSeverity = "warning"


class PersonalDetails(BaseModel):
    """Normalized personal and contact details extracted from a CV."""

    model_config = ConfigDict(extra="forbid")

    full_name: str | None = None
    headline: str | None = None
    emails: list[str] = Field(default_factory=list)
    phones: list[str] = Field(default_factory=list)
    location: str | None = None
    linkedin_url: str | None = None
    github_url: str | None = None
    portfolio_urls: list[str] = Field(default_factory=list)


class ExperienceEntry(BaseModel):
    """Normalized professional experience record."""

    model_config = ConfigDict(extra="forbid")

    title: str | None = None
    company_name: str | None = None
    location: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    is_current: bool = False
    summary: str | None = None
    bullets: list[str] = Field(default_factory=list)


class EducationEntry(BaseModel):
    """Normalized education record."""

    model_config = ConfigDict(extra="forbid")

    school_name: str | None = None
    degree_name: str | None = None
    field_of_study: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    grade: str | None = None


class LanguageEntry(BaseModel):
    """Normalized language skill record."""

    model_config = ConfigDict(extra="forbid")

    language_name: str
    proficiency_level: str | None = None
    notes: str | None = None


class CertificationEntry(BaseModel):
    """Normalized certification or license record."""

    model_config = ConfigDict(extra="forbid")

    name: str
    issuer: str | None = None
    issued_date: str | None = None
    credential_id: str | None = None


class CVExtractionMetadata(BaseModel):
    """Extraction metadata and quality telemetry for a parsed document."""

    model_config = ConfigDict(extra="forbid")

    status: ParseStatus
    file_type: DocumentFileType
    mime_type: str | None = None
    extraction_method: str
    used_ocr: bool = False
    ocr_backend: str | None = None
    page_count: int = 0
    text_char_count: int = 0
    text_line_count: int = 0
    quality_score: int = 0
    quality_label: QualityLabel = "low"
    extraction_duration_ms: int = 0
    parsing_duration_ms: int = 0
    total_duration_ms: int = 0
    warnings: list[ParseWarning] = Field(default_factory=list)


class CVStructuredProfile(BaseModel):
    """Canonical structured representation of a parsed CV."""

    model_config = ConfigDict(extra="forbid")

    personal_details: PersonalDetails = Field(default_factory=PersonalDetails)
    summary: str | None = None
    skills: list[str] = Field(default_factory=list)
    target_roles: list[str] = Field(default_factory=list)
    preferred_locations: list[str] = Field(default_factory=list)
    remote_preference: str | None = None
    experience_entries: list[ExperienceEntry] = Field(default_factory=list)
    education_entries: list[EducationEntry] = Field(default_factory=list)
    language_entries: list[LanguageEntry] = Field(default_factory=list)
    certifications: list[CertificationEntry] = Field(default_factory=list)


class CVParseResult(BaseModel):
    """Top-level parse output for a single input document."""

    model_config = ConfigDict(extra="forbid")

    parser_version: str = PARSER_VERSION
    input_filename: str
    source_path: str | None = None
    content_sha256: str
    parsed_at: str
    metadata: CVExtractionMetadata
    profile: CVStructuredProfile = Field(default_factory=CVStructuredProfile)
    section_map: dict[str, str] = Field(default_factory=dict)
    raw_text: str | None = None
    normalized_text: str | None = None


@dataclass(slots=True)
class ExtractionArtifact:
    """Internal result of the extraction stage."""

    file_type: DocumentFileType
    mime_type: str | None
    raw_text: str
    normalized_text: str
    extraction_method: str
    used_ocr: bool
    ocr_backend: str | None
    page_count: int
    warnings: list[ParseWarning] = field(default_factory=list)


class CVParserEngineError(RuntimeError):
    """Base error class for parser engine failures."""


class UnsupportedDocumentError(CVParserEngineError):
    """Raised when an input document type is unsupported."""


class OCRBackendUnavailableError(CVParserEngineError):
    """Raised when OCR is required but no OCR backend is available."""


class UniversalCVParserEngine:
    """Enterprise-grade parser for CVs across PDF, DOCX, and image formats.

    The engine auto-detects the input type, selects an extraction strategy,
    falls back to OCR when necessary, and normalizes extracted content into a
    Pydantic-backed structured output model.
    """

    def __init__(self, config: ParserConfig | None = None) -> None:
        self.config = config or ParserConfig()
        self._easyocr_reader_cache: dict[tuple[str, ...], Any] = {}

    def parse_file(self, file_path: str | os.PathLike[str]) -> CVParseResult:
        """Parse a single CV file into a structured, normalized output model."""
        started_at = time.perf_counter()
        path = Path(file_path)
        filename = path.name
        sha256 = self._sha256_for_file(path)
        parsed_at = self._utc_now_iso()
        warnings: list[ParseWarning] = []

        try:
            mime_type = self._detect_mime_type(path)
            extraction_started_at = time.perf_counter()
            artifact = self._extract_text(path, mime_type)
            extraction_duration_ms = self._elapsed_ms(extraction_started_at)

            parsing_started_at = time.perf_counter()
            profile, section_map, parsing_warnings = self._parse_structured_profile(
                artifact.normalized_text
            )
            parsing_duration_ms = self._elapsed_ms(parsing_started_at)
            warnings.extend(artifact.warnings)
            warnings.extend(parsing_warnings)

            quality_score = self._compute_quality_score(
                text=artifact.normalized_text,
                page_count=artifact.page_count,
                used_ocr=artifact.used_ocr,
                profile=profile,
            )
            quality_label = self._quality_label_from_score(quality_score)
            status = self._derive_status(
                normalized_text=artifact.normalized_text,
                warnings=warnings,
                profile=profile,
            )

            metadata = CVExtractionMetadata(
                status=status,
                file_type=artifact.file_type,
                mime_type=artifact.mime_type,
                extraction_method=artifact.extraction_method,
                used_ocr=artifact.used_ocr,
                ocr_backend=artifact.ocr_backend,
                page_count=artifact.page_count,
                text_char_count=len(artifact.normalized_text),
                text_line_count=self._line_count(artifact.normalized_text),
                quality_score=quality_score,
                quality_label=quality_label,
                extraction_duration_ms=extraction_duration_ms,
                parsing_duration_ms=parsing_duration_ms,
                total_duration_ms=self._elapsed_ms(started_at),
                warnings=warnings,
            )

            return CVParseResult(
                input_filename=filename,
                source_path=str(path),
                content_sha256=sha256,
                parsed_at=parsed_at,
                metadata=metadata,
                profile=profile,
                section_map=section_map,
                raw_text=(artifact.raw_text if self.config.include_raw_text else None),
                normalized_text=(
                    artifact.normalized_text
                    if self.config.include_normalized_text
                    else None
                ),
            )
        except Exception as exc:  # pragma: no cover - production safety net
            LOGGER.exception("CV parsing failed for %s", path)
            warnings.append(
                ParseWarning(
                    code="parser_failure",
                    message=f"Parsing failed: {exc}",
                    severity="error",
                )
            )
            metadata = CVExtractionMetadata(
                status="failed",
                file_type=self._detect_file_type(path, None),
                mime_type=self._detect_mime_type(path),
                extraction_method="failed",
                total_duration_ms=self._elapsed_ms(started_at),
                warnings=warnings,
            )
            return CVParseResult(
                input_filename=filename,
                source_path=str(path),
                content_sha256=sha256,
                parsed_at=parsed_at,
                metadata=metadata,
                profile=CVStructuredProfile(),
                section_map={},
                raw_text=None,
                normalized_text=None,
            )

    def parse_bytes(
        self,
        content: bytes,
        *,
        filename: str,
        content_type: str | None = None,
    ) -> CVParseResult:
        """Parse a document from bytes using the same resilient pipeline."""
        suffix = Path(filename).suffix or self._suffix_from_content_type(content_type)
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
            handle.write(content)
            temp_path = Path(handle.name)
        try:
            result = self.parse_file(temp_path)
            result.input_filename = filename
            result.source_path = None
            if content_type:
                result.metadata.mime_type = content_type
            result.content_sha256 = hashlib.sha256(content).hexdigest()
            return result
        finally:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:  # pragma: no cover - cleanup best effort
                LOGGER.warning("Failed to remove temporary file %s", temp_path)

    async def parse_file_async(
        self,
        file_path: str | os.PathLike[str],
    ) -> CVParseResult:
        """Asynchronously parse a single file using a thread offload."""
        return await asyncio.to_thread(self.parse_file, file_path)

    async def parse_many_async(
        self,
        file_paths: Sequence[str | os.PathLike[str]],
        *,
        concurrency: int | None = None,
    ) -> list[CVParseResult]:
        """Asynchronously parse multiple CVs with bounded concurrency."""
        semaphore = asyncio.Semaphore(concurrency or self.config.async_concurrency)

        async def _parse_one(path: str | os.PathLike[str]) -> CVParseResult:
            async with semaphore:
                return await self.parse_file_async(path)

        return await asyncio.gather(*[_parse_one(path) for path in file_paths])

    def parse_many_multiprocess(
        self,
        file_paths: Sequence[str | os.PathLike[str]],
        *,
        max_workers: int | None = None,
    ) -> list[CVParseResult]:
        """Parse multiple documents with process-level parallelism."""
        worker_count = max_workers or self.config.process_workers
        config_dict = self.config.model_dump(mode="json")
        with ProcessPoolExecutor(max_workers=worker_count) as executor:
            futures = [
                executor.submit(_parse_file_worker, str(Path(path)), config_dict)
                for path in file_paths
            ]
            return [future.result() for future in futures]

    def _extract_text(
        self,
        path: Path,
        mime_type: str | None,
    ) -> ExtractionArtifact:
        file_type = self._detect_file_type(path, mime_type)
        if file_type == "pdf":
            return self._extract_text_from_pdf(path, mime_type)
        if file_type == "docx":
            return self._extract_text_from_docx(path, mime_type)
        if file_type in {"png", "jpg", "jpeg", "image"}:
            return self._extract_text_from_image(path, mime_type, file_type)
        raise UnsupportedDocumentError(
            f"Unsupported CV file type for '{path.name}'."
        )

    def _extract_text_from_pdf(
        self,
        path: Path,
        mime_type: str | None,
    ) -> ExtractionArtifact:
        warnings: list[ParseWarning] = []
        page_texts: list[str] = []
        used_ocr = False
        ocr_backend: str | None = None

        if fitz is not None:
            try:
                with fitz.open(path) as document:
                    if getattr(document, "needs_pass", False):
                        warnings.append(
                            ParseWarning(
                                code="encrypted_pdf",
                                message=(
                                    "PDF appears encrypted or password protected; "
                                    "text extraction may be incomplete."
                                ),
                                severity="warning",
                            )
                        )
                    for page in document:
                        page_text = self._normalize_text(
                            page.get_text("text", sort=True)
                        )
                        if len(page_text) < self.config.pdf_sparse_text_threshold:
                            try:
                                image = self._render_pdf_page_to_image(page)
                                ocr_text, backend_name, ocr_warnings = self._ocr_image(
                                    image
                                )
                                warnings.extend(ocr_warnings)
                                if len(ocr_text) > len(page_text):
                                    page_text = ocr_text
                                    used_ocr = True
                                    ocr_backend = backend_name or ocr_backend
                            except Exception as exc:
                                warnings.append(
                                    ParseWarning(
                                        code="pdf_page_ocr_failed",
                                        message=(
                                            f"OCR fallback failed for page "
                                            f"{page.number + 1}: {exc}"
                                        ),
                                        severity="warning",
                                    )
                                )
                        page_texts.append(page_text)
                normalized_text = self._normalize_text("\n\n".join(page_texts))
                return ExtractionArtifact(
                    file_type="pdf",
                    mime_type=mime_type,
                    raw_text="\n\n".join(page_texts),
                    normalized_text=normalized_text,
                    extraction_method="pymupdf",
                    used_ocr=used_ocr,
                    ocr_backend=ocr_backend,
                    page_count=len(page_texts),
                    warnings=warnings,
                )
            except Exception as exc:
                warnings.append(
                    ParseWarning(
                        code="pymupdf_failed",
                        message=f"PyMuPDF extraction failed: {exc}",
                        severity="warning",
                    )
                )
                LOGGER.warning("PyMuPDF failed for %s: %s", path, exc)
        else:
            warnings.append(
                ParseWarning(
                    code="pymupdf_missing",
                    message="PyMuPDF is not installed; skipping primary PDF reader.",
                    severity="info",
                )
            )

        if pdfplumber is not None:
            try:
                with pdfplumber.open(path) as pdf_document:
                    for page in pdf_document.pages:
                        page_texts.append(
                            self._normalize_text(page.extract_text() or "")
                        )
                normalized_text = self._normalize_text("\n\n".join(page_texts))
                return ExtractionArtifact(
                    file_type="pdf",
                    mime_type=mime_type,
                    raw_text="\n\n".join(page_texts),
                    normalized_text=normalized_text,
                    extraction_method="pdfplumber",
                    used_ocr=False,
                    ocr_backend=None,
                    page_count=len(page_texts),
                    warnings=warnings,
                )
            except Exception as exc:
                warnings.append(
                    ParseWarning(
                        code="pdfplumber_failed",
                        message=f"pdfplumber extraction failed: {exc}",
                        severity="warning",
                    )
                )
                LOGGER.warning("pdfplumber failed for %s: %s", path, exc)
        else:
            warnings.append(
                ParseWarning(
                    code="pdfplumber_missing",
                    message="pdfplumber is not installed; PDF fallback reader unavailable.",
                    severity="info",
                )
            )

        raise CVParserEngineError(
            "PDF extraction failed with all configured extraction backends."
        )

    def _extract_text_from_docx(
        self,
        path: Path,
        mime_type: str | None,
    ) -> ExtractionArtifact:
        if Document is None:
            raise CVParserEngineError(
                "python-docx is not installed; DOCX parsing is unavailable."
            )

        warnings: list[ParseWarning] = []
        document = Document(path)
        paragraphs = [
            self._normalize_text(paragraph.text)
            for paragraph in document.paragraphs
            if self._normalize_text(paragraph.text)
        ]

        table_lines: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [self._normalize_text(cell.text) for cell in row.cells]
                filtered_cells = [cell for cell in cells if cell]
                if filtered_cells:
                    table_lines.append(" | ".join(filtered_cells))

        raw_text = "\n".join([*paragraphs, *table_lines])
        normalized_text = self._normalize_text(raw_text)
        if not normalized_text:
            warnings.append(
                ParseWarning(
                    code="docx_empty",
                    message="DOCX parsing succeeded but produced empty text.",
                    severity="warning",
                )
            )

        return ExtractionArtifact(
            file_type="docx",
            mime_type=mime_type,
            raw_text=raw_text,
            normalized_text=normalized_text,
            extraction_method="python-docx",
            used_ocr=False,
            ocr_backend=None,
            page_count=max(1, len(document.sections)),
            warnings=warnings,
        )

    def _extract_text_from_image(
        self,
        path: Path,
        mime_type: str | None,
        file_type: DocumentFileType,
    ) -> ExtractionArtifact:
        if Image is None:
            raise CVParserEngineError(
                "Pillow is not installed; image-based CV parsing is unavailable."
            )

        warnings: list[ParseWarning] = []
        with Image.open(path) as source_image:
            image = source_image.convert("RGB")
            preprocessed = self._preprocess_image_for_ocr(image)
            text, backend_name, ocr_warnings = self._ocr_image(preprocessed)
            warnings.extend(ocr_warnings)

        return ExtractionArtifact(
            file_type=file_type,
            mime_type=mime_type,
            raw_text=text,
            normalized_text=self._normalize_text(text),
            extraction_method="ocr-image",
            used_ocr=True,
            ocr_backend=backend_name,
            page_count=1,
            warnings=warnings,
        )

    def _ocr_image(self, image: Any) -> tuple[str, str | None, list[ParseWarning]]:
        warnings: list[ParseWarning] = []
        preferred_backends = self.config.preferred_ocr_backends

        for backend_name in preferred_backends:
            if backend_name == "pytesseract":
                if pytesseract is None:
                    warnings.append(
                        ParseWarning(
                            code="pytesseract_missing",
                            message="pytesseract is not installed.",
                            severity="info",
                        )
                    )
                    continue
                try:
                    text = pytesseract.image_to_string(
                        image,
                        lang="+".join(self.config.ocr_languages),
                        config="--oem 3 --psm 6",
                    )
                    return self._normalize_text(text), "pytesseract", warnings
                except Exception as exc:
                    warnings.append(
                        ParseWarning(
                            code="pytesseract_failed",
                            message=f"pytesseract OCR failed: {exc}",
                            severity="warning",
                        )
                    )
                    LOGGER.warning("pytesseract OCR failed: %s", exc)
                    continue

            if backend_name == "easyocr":
                if easyocr is None:
                    warnings.append(
                        ParseWarning(
                            code="easyocr_missing",
                            message="EasyOCR is not installed.",
                            severity="info",
                        )
                    )
                    continue
                try:
                    reader = self._get_easyocr_reader(self.config.ocr_languages)
                    raw_result = reader.readtext(image, detail=0, paragraph=True)
                    if isinstance(raw_result, list):
                        text = "\n".join(str(item) for item in raw_result if item)
                    else:
                        text = str(raw_result)
                    return self._normalize_text(text), "easyocr", warnings
                except Exception as exc:
                    warnings.append(
                        ParseWarning(
                            code="easyocr_failed",
                            message=f"EasyOCR failed: {exc}",
                            severity="warning",
                        )
                    )
                    LOGGER.warning("EasyOCR failed: %s", exc)
                    continue

        raise OCRBackendUnavailableError(
            "No OCR backend is available. Install pytesseract and/or EasyOCR."
        )

    def _get_easyocr_reader(self, languages: tuple[str, ...]) -> Any:
        if languages in self._easyocr_reader_cache:
            return self._easyocr_reader_cache[languages]
        if easyocr is None:
            raise OCRBackendUnavailableError("EasyOCR is not installed.")
        reader = easyocr.Reader(list(languages), gpu=False)
        self._easyocr_reader_cache[languages] = reader
        return reader

    def _preprocess_image_for_ocr(self, image: Any) -> Any:
        if ImageOps is None or ImageFilter is None:
            return image

        grayscale = ImageOps.grayscale(image)
        autocontrasted = ImageOps.autocontrast(grayscale)
        sharpened = autocontrasted.filter(ImageFilter.SHARPEN)
        median = sharpened.filter(ImageFilter.MedianFilter(size=3))
        thresholded = median.point(lambda value: 255 if value > 165 else 0)
        return thresholded.convert("RGB")

    def _render_pdf_page_to_image(self, page: Any) -> Any:
        if fitz is None or Image is None:
            raise CVParserEngineError(
                "PDF page rendering requires both PyMuPDF and Pillow."
            )
        matrix = fitz.Matrix(self.config.render_scale, self.config.render_scale)
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        return Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")

    def _parse_structured_profile(
        self,
        normalized_text: str,
    ) -> tuple[CVStructuredProfile, dict[str, str], list[ParseWarning]]:
        warnings: list[ParseWarning] = []
        if not normalized_text:
            warnings.append(
                ParseWarning(
                    code="empty_text",
                    message="No readable text was extracted from the document.",
                    severity="warning",
                )
            )
            return CVStructuredProfile(), {}, warnings

        preamble_lines, sections = self._segment_sections(normalized_text)
        personal_details = self._extract_personal_details(preamble_lines)
        summary = self._extract_summary(preamble_lines, sections, personal_details)
        skills = self._extract_skills(sections, normalized_text)
        experience_entries = self._extract_experience_entries(sections, normalized_text)
        education_entries = self._extract_education_entries(sections, normalized_text)
        language_entries = self._extract_language_entries(sections, normalized_text)
        certifications = self._extract_certifications(sections)
        target_roles = self._extract_target_roles(
            sections=sections,
            personal_details=personal_details,
            experience_entries=experience_entries,
        )
        preferred_locations = self._extract_preferred_locations(
            preamble_lines,
            sections,
            personal_details,
        )
        remote_preference = self._extract_remote_preference(
            normalized_text,
            sections,
        )

        if not personal_details.headline and summary:
            headline_from_summary = self._derive_headline_from_summary(summary)
            if headline_from_summary:
                personal_details.headline = headline_from_summary

        if not skills and len(normalized_text) >= 200:
            warnings.append(
                ParseWarning(
                    code="skills_not_found",
                    message=(
                        "No skills section could be confidently identified; skill "
                        "coverage may be incomplete."
                    ),
                    severity="info",
                )
            )
        if not experience_entries:
            warnings.append(
                ParseWarning(
                    code="experience_not_found",
                    message=(
                        "No experience entries were confidently extracted; "
                        "review the structured output before applying changes."
                    ),
                    severity="info",
                )
            )

        section_map = {
            section_name: self._normalize_text("\n".join(lines))
            for section_name, lines in sections.items()
            if self._normalize_text("\n".join(lines))
        }

        profile = CVStructuredProfile(
            personal_details=personal_details,
            summary=summary,
            skills=skills,
            target_roles=target_roles,
            preferred_locations=preferred_locations,
            remote_preference=remote_preference,
            experience_entries=experience_entries,
            education_entries=education_entries,
            language_entries=language_entries,
            certifications=certifications,
        )
        return profile, section_map, warnings

    def _segment_sections(
        self,
        normalized_text: str,
    ) -> tuple[list[str], dict[str, list[str]]]:
        lines = [line.strip() for line in normalized_text.split("\n")]
        preamble: list[str] = []
        sections: dict[str, list[str]] = {}
        current_section: str | None = None

        for line in lines:
            detected = self._detect_section_name(line)
            if detected is not None:
                current_section = detected
                sections.setdefault(current_section, [])
                continue
            if current_section is None:
                preamble.append(line)
            else:
                sections[current_section].append(line)

        return preamble, sections

    def _detect_section_name(self, line: str) -> str | None:
        cleaned = line.strip().rstrip(":")
        if not cleaned:
            return None
        if EMAIL_RE.search(cleaned) or PHONE_RE.search(cleaned) or URL_RE.search(cleaned):
            return None
        if len(cleaned.split()) > SECTION_NAME_MAX_WORDS:
            return None
        normalized = _normalize_alias_key(cleaned)
        if not normalized:
            return None

        for canonical_name, aliases in NORMALIZED_SECTION_ALIASES.items():
            if normalized in aliases:
                return canonical_name

        return None

    def _extract_personal_details(self, preamble_lines: list[str]) -> PersonalDetails:
        meaningful_lines = [line for line in preamble_lines if line][:CONTACT_SCAN_LINE_LIMIT]
        full_name = self._extract_name_candidate(meaningful_lines)
        headline = self._extract_headline_candidate(meaningful_lines, full_name)
        emails = self._extract_unique_matches(EMAIL_RE, meaningful_lines)
        phones = self._normalize_phones(
            self._extract_unique_matches(PHONE_RE, meaningful_lines)
        )
        urls = self._extract_unique_matches(URL_RE, meaningful_lines)
        linkedin_url = next((url for url in urls if LINKEDIN_RE.search(url)), None)
        github_url = next((url for url in urls if GITHUB_RE.search(url)), None)
        portfolio_urls = [
            url
            for url in urls
            if url not in {linkedin_url, github_url}
        ][: self.config.max_urls]
        location = self._extract_location_candidate(meaningful_lines)

        return PersonalDetails(
            full_name=full_name,
            headline=headline,
            emails=emails,
            phones=phones,
            location=location,
            linkedin_url=linkedin_url,
            github_url=github_url,
            portfolio_urls=portfolio_urls,
        )

    def _extract_summary(
        self,
        preamble_lines: list[str],
        sections: dict[str, list[str]],
        personal_details: PersonalDetails,
    ) -> str | None:
        summary_lines = [
            line for line in sections.get("summary", []) if line and not self._looks_like_contact_line(line)
        ]
        if summary_lines:
            return self._truncate_text(" ".join(summary_lines), self.config.max_summary_chars)

        candidates: list[str] = []
        for line in preamble_lines:
            if not line:
                if candidates:
                    break
                continue
            if line == personal_details.full_name:
                continue
            if line == personal_details.headline:
                continue
            if self._looks_like_contact_line(line):
                continue
            if self._looks_like_section_heading_candidate(line):
                break
            candidates.append(line)
            if len(candidates) >= SUMMARY_SCAN_LINE_LIMIT:
                break

        if not candidates:
            return None
        return self._truncate_text(" ".join(candidates), self.config.max_summary_chars)

    def _extract_skills(
        self,
        sections: dict[str, list[str]],
        normalized_text: str,
    ) -> list[str]:
        skills: list[str] = []
        section_lines = sections.get("skills", [])
        for line in section_lines:
            skills.extend(self._extract_skill_tokens_from_line(line))

        if not skills:
            lowered_text = normalized_text.casefold()
            for skill in SKILL_LEXICON:
                if skill in lowered_text:
                    skills.append(self._display_skill(skill))

        return self._deduplicate_strings(skills)[: self.config.max_skills]

    def _extract_experience_entries(
        self,
        sections: dict[str, list[str]],
        normalized_text: str,
    ) -> list[ExperienceEntry]:
        source_lines = sections.get("experience", [])
        if not any(line for line in source_lines):
            source_lines = self._fallback_experience_lines(normalized_text)

        blocks = self._split_blocks(source_lines)
        entries: list[ExperienceEntry] = []
        for block in blocks[:MAX_BLOCKS_PER_SECTION]:
            entry = self._parse_experience_block(block)
            if entry is not None:
                entries.append(entry)
        return entries

    def _extract_education_entries(
        self,
        sections: dict[str, list[str]],
        normalized_text: str,
    ) -> list[EducationEntry]:
        source_lines = sections.get("education", [])
        if not any(line for line in source_lines):
            source_lines = self._fallback_education_lines(normalized_text)

        blocks = self._split_blocks(source_lines)
        entries: list[EducationEntry] = []
        for block in blocks[:MAX_BLOCKS_PER_SECTION]:
            entry = self._parse_education_block(block)
            if entry is not None:
                entries.append(entry)
        return entries

    def _extract_language_entries(
        self,
        sections: dict[str, list[str]],
        normalized_text: str,
    ) -> list[LanguageEntry]:
        section_lines = [line for line in sections.get("languages", []) if line]
        if not section_lines:
            section_lines = self._fallback_language_lines(normalized_text)

        entries: list[LanguageEntry] = []
        for line in section_lines:
            entry = self._parse_language_line(line)
            if entry is not None:
                entries.append(entry)
        return entries

    def _extract_certifications(
        self,
        sections: dict[str, list[str]],
    ) -> list[CertificationEntry]:
        lines = [line for line in sections.get("certifications", []) if line]
        blocks = self._split_blocks(lines)
        certifications: list[CertificationEntry] = []
        for block in blocks[: self.config.max_certifications]:
            certification = self._parse_certification_block(block)
            if certification is not None:
                certifications.append(certification)
        return certifications

    def _extract_target_roles(
        self,
        *,
        sections: dict[str, list[str]],
        personal_details: PersonalDetails,
        experience_entries: list[ExperienceEntry],
    ) -> list[str]:
        explicit_roles: list[str] = []
        for line in sections.get("target_roles", []):
            explicit_roles.extend(self._split_candidate_tokens(line))
        if explicit_roles:
            return self._deduplicate_strings(
                [role for role in explicit_roles if self._looks_like_role_line(role)]
            )[: self.config.max_target_roles]

        roles: list[str] = []
        if personal_details.headline and self._looks_like_role_line(
            personal_details.headline
        ):
            roles.append(personal_details.headline)
        for entry in experience_entries[:5]:
            if entry.title and self._looks_like_role_line(entry.title):
                roles.append(entry.title)
        return self._deduplicate_strings(roles)[: self.config.max_target_roles]

    def _extract_preferred_locations(
        self,
        preamble_lines: list[str],
        sections: dict[str, list[str]],
        personal_details: PersonalDetails,
    ) -> list[str]:
        candidates: list[str] = []
        for line in sections.get("preferences", []):
            candidates.extend(self._extract_location_tokens(line))

        if personal_details.location:
            candidates.append(personal_details.location)

        if not candidates:
            for line in preamble_lines[:CONTACT_SCAN_LINE_LIMIT]:
                candidates.extend(self._extract_location_tokens(line))

        return self._deduplicate_strings(candidates)

    def _extract_remote_preference(
        self,
        normalized_text: str,
        sections: dict[str, list[str]],
    ) -> str | None:
        source = "\n".join(sections.get("preferences", [])) or normalized_text
        lowered = source.casefold()
        if "hybrid" in lowered or "hibrit" in lowered:
            return "hybrid"
        if "remote" in lowered or "uzaktan" in lowered:
            return "remote"
        if "onsite" in lowered or "on-site" in lowered or "vor ort" in lowered:
            return "onsite"
        return None

    def _parse_experience_block(self, block: list[str]) -> ExperienceEntry | None:
        content_lines = [line for line in block if line]
        if not content_lines:
            return None

        date_line_index: int | None = None
        start_date: str | None = None
        end_date: str | None = None
        is_current = False

        for index, line in enumerate(content_lines[:4]):
            start_date, end_date, is_current = self._extract_date_range(line)
            if start_date or end_date:
                date_line_index = index
                break

        header_candidates = [
            line
            for line in content_lines[:3]
            if not BULLET_PREFIX_RE.match(line)
        ]
        title, company_name, location = self._extract_title_company_location(
            header_candidates,
            date_line_index,
        )

        summary_lines: list[str] = []
        bullets: list[str] = []
        for index, line in enumerate(content_lines):
            if index == date_line_index:
                continue
            if line in {title, company_name}:
                continue
            if BULLET_PREFIX_RE.match(line):
                bullets.append(BULLET_PREFIX_RE.sub("", line).strip())
            elif index > 0:
                summary_lines.append(line)

        summary = self._normalize_text(" ".join(summary_lines[:8]))
        if not title and not company_name:
            return None

        return ExperienceEntry(
            title=title,
            company_name=company_name,
            location=location,
            start_date=start_date,
            end_date=end_date,
            is_current=is_current,
            summary=summary,
            bullets=bullets[:MAX_BULLETS_PER_ENTRY],
        )

    def _parse_education_block(self, block: list[str]) -> EducationEntry | None:
        content_lines = [line for line in block if line]
        if not content_lines:
            return None

        start_date: str | None = None
        end_date: str | None = None
        for line in content_lines[:3]:
            start_date, end_date, _ = self._extract_date_range(line)
            if start_date or end_date:
                break

        school_name: str | None = None
        degree_name: str | None = None
        field_of_study: str | None = None

        for line in content_lines[:4]:
            lowered = line.casefold()
            if school_name is None and self._looks_like_school_line(lowered):
                school_name = line
                continue
            if degree_name is None and self._looks_like_degree_line(lowered):
                degree_name = line
                continue
            if field_of_study is None and self._looks_like_field_of_study_line(lowered):
                field_of_study = line

        if school_name is None and content_lines:
            school_name = content_lines[0]
        if degree_name is None and len(content_lines) > 1:
            degree_name = content_lines[1]

        if school_name is None:
            return None

        grade = None
        for line in content_lines:
            if "gpa" in line.casefold() or "grade" in line.casefold():
                grade = line
                break

        return EducationEntry(
            school_name=school_name,
            degree_name=degree_name,
            field_of_study=field_of_study,
            start_date=start_date,
            end_date=end_date,
            grade=grade,
        )

    def _parse_language_line(self, line: str) -> LanguageEntry | None:
        clean_line = self._normalize_text(line)
        if not clean_line:
            return None
        tokens = [token.strip() for token in PIPE_SPLIT_RE.split(clean_line) if token.strip()]
        if len(tokens) == 1:
            tokens = [token.strip() for token in COMMA_SPLIT_RE.split(clean_line) if token.strip()]
        if len(tokens) == 1:
            tokens = [token.strip() for token in SEMICOLON_SPLIT_RE.split(clean_line) if token.strip()]

        language_name: str | None = None
        proficiency_level: str | None = None
        notes: str | None = None

        for token in tokens:
            lowered = token.casefold()
            if language_name is None and any(
                name in lowered for name in KNOWN_LANGUAGE_NAMES
            ):
                language_name = token
                continue
            if proficiency_level is None and any(
                hint in lowered for hint in LANGUAGE_PROFICIENCY_HINTS
            ):
                proficiency_level = token
                continue
        if language_name is None:
            first = tokens[0] if tokens else clean_line
            first_word = first.split()[0].casefold()
            if first_word not in KNOWN_LANGUAGE_NAMES:
                return None
            language_name = first

        if len(tokens) > 1:
            residual = [
                token
                for token in tokens
                if token != language_name and token != proficiency_level
            ]
            if residual:
                notes = "; ".join(residual)

        return LanguageEntry(
            language_name=language_name,
            proficiency_level=proficiency_level,
            notes=notes,
        )

    def _parse_certification_block(
        self,
        block: list[str],
    ) -> CertificationEntry | None:
        content_lines = [line for line in block if line]
        if not content_lines:
            return None
        name = content_lines[0]
        issuer = content_lines[1] if len(content_lines) > 1 else None
        issued_date = None
        credential_id = None
        for line in content_lines[1:]:
            if YEAR_RE.search(line) or DATE_RANGE_RE.search(line):
                issued_date = line
            if "credential" in line.casefold() or "id" in line.casefold():
                credential_id = line
        return CertificationEntry(
            name=name,
            issuer=issuer,
            issued_date=issued_date,
            credential_id=credential_id,
        )

    def _fallback_experience_lines(self, normalized_text: str) -> list[str]:
        lines = [line.strip() for line in normalized_text.split("\n") if line.strip()]
        return [line for line in lines if self._looks_like_experience_line(line)]

    def _fallback_education_lines(self, normalized_text: str) -> list[str]:
        lines = [line.strip() for line in normalized_text.split("\n") if line.strip()]
        return [line for line in lines if self._looks_like_education_line(line)]

    def _fallback_language_lines(self, normalized_text: str) -> list[str]:
        lines = [line.strip() for line in normalized_text.split("\n") if line.strip()]
        return [line for line in lines if self._looks_like_language_line(line)]

    def _extract_name_candidate(self, lines: list[str]) -> str | None:
        for line in lines[:6]:
            if not line:
                continue
            if self._looks_like_contact_line(line):
                continue
            if self._looks_like_section_heading_candidate(line):
                continue
            if YEAR_RE.search(line):
                continue
            tokens = line.split()
            if len(tokens) < 2 or len(tokens) > MAX_NAME_TOKENS:
                continue
            if any(token.isdigit() for token in tokens):
                continue
            if sum(1 for token in tokens if token[:1].isupper()) >= max(2, len(tokens) - 1):
                return line
        return None

    def _extract_headline_candidate(
        self,
        lines: list[str],
        full_name: str | None,
    ) -> str | None:
        for line in lines[:8]:
            if not line:
                continue
            if line == full_name:
                continue
            if self._looks_like_contact_line(line):
                continue
            if len(line) > MAX_HEADLINE_LENGTH:
                continue
            if self._looks_like_role_line(line):
                return line
        return None

    def _extract_location_candidate(self, lines: list[str]) -> str | None:
        for line in lines[:8]:
            if not line:
                continue
            if EMAIL_RE.search(line) or PHONE_RE.search(line):
                parts = [part.strip() for part in PIPE_SPLIT_RE.split(line)]
                for part in parts:
                    if not part:
                        continue
                    if EMAIL_RE.search(part) or PHONE_RE.search(part):
                        continue
                    if URL_RE.search(part):
                        continue
                    if len(part.split()) <= 4 and any(char.isalpha() for char in part):
                        return part
            if "," in line and not URL_RE.search(line):
                if len(line.split()) <= 6:
                    return line
        return None

    def _extract_location_tokens(self, line: str) -> list[str]:
        if not line or URL_RE.search(line):
            return []
        tokens = [token.strip() for token in PIPE_SPLIT_RE.split(line) if token.strip()]
        locations = [
            token
            for token in tokens
            if not EMAIL_RE.search(token)
            and not PHONE_RE.search(token)
            and not any(keyword in token.casefold() for keyword in REMOTE_KEYWORDS)
            and 1 <= len(token.split()) <= 4
        ]
        return locations

    def _extract_skill_tokens_from_line(self, line: str) -> list[str]:
        clean_line = self._normalize_text(line)
        if not clean_line or self._looks_like_contact_line(clean_line):
            return []

        tokens = self._split_candidate_tokens(clean_line)
        skills: list[str] = []
        for token in tokens:
            lowered = token.casefold()
            if lowered in {"skills", "technical skills", "technologies"}:
                continue
            if lowered in KNOWN_LANGUAGE_NAMES:
                continue
            if self._looks_like_company_line(lowered) or self._looks_like_school_line(lowered):
                continue
            if lowered in SKILL_LEXICON or len(token) <= 32:
                skills.append(self._display_skill(token))
        return skills

    def _extract_date_range(self, line: str) -> tuple[str | None, str | None, bool]:
        match = DATE_RANGE_RE.search(line)
        if not match:
            years = YEAR_RE.findall(line)
            if len(years) >= 2:
                return years[0], years[1], False
            return None, None, False
        start_date = match.group("start")
        end_date = match.group("end")
        is_current = end_date.casefold() in {
            "present",
            "current",
            "ongoing",
            "today",
            "heute",
            "devam",
        }
        return start_date, end_date, is_current

    def _extract_title_company_location(
        self,
        header_candidates: list[str],
        date_line_index: int | None,
    ) -> tuple[str | None, str | None, str | None]:
        header_lines = [line for line in header_candidates if line]
        if date_line_index is not None and date_line_index < len(header_lines):
            header_lines = [
                line
                for index, line in enumerate(header_lines)
                if index != date_line_index
            ]

        title: str | None = None
        company_name: str | None = None
        location: str | None = None

        for line in header_lines:
            if " at " in line.casefold():
                left, right = re.split(r"\s+at\s+", line, maxsplit=1, flags=re.IGNORECASE)
                title = left.strip() or title
                company_name = right.strip() or company_name
                break
            if " @ " in line:
                left, right = line.split(" @ ", maxsplit=1)
                title = left.strip() or title
                company_name = right.strip() or company_name
                break
            split_by_pipe = [part.strip() for part in PIPE_SPLIT_RE.split(line) if part.strip()]
            if len(split_by_pipe) >= 2 and title is None and company_name is None:
                if self._looks_like_role_line(split_by_pipe[0]):
                    title = split_by_pipe[0]
                    company_name = split_by_pipe[1]
                    if len(split_by_pipe) >= 3:
                        location = split_by_pipe[2]
                    break

        if title is None and header_lines:
            for candidate in header_lines:
                if self._looks_like_role_line(candidate):
                    title = candidate
                    break
        if company_name is None and len(header_lines) >= 2:
            for candidate in header_lines:
                if candidate != title and self._looks_like_company_line(candidate.casefold()):
                    company_name = candidate
                    break
        if company_name is None and len(header_lines) >= 2:
            company_name = next((line for line in header_lines if line != title), None)

        if location is None:
            for line in header_lines:
                if line not in {title, company_name}:
                    possible_locations = self._extract_location_tokens(line)
                    if possible_locations:
                        location = possible_locations[0]
                        break

        return title, company_name, location

    def _split_blocks(self, lines: Sequence[str]) -> list[list[str]]:
        blocks: list[list[str]] = []
        current: list[str] = []

        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                if current:
                    blocks.append(current)
                    current = []
                continue
            if current and self._should_start_new_block(line, current):
                blocks.append(current)
                current = [line]
                continue
            current.append(line)

        if current:
            blocks.append(current)
        return blocks

    def _should_start_new_block(self, line: str, current_block: list[str]) -> bool:
        if len(current_block) >= 8:
            return True
        if DATE_RANGE_RE.search(line) and current_block:
            return True
        if self._looks_like_role_line(line) and current_block and not BULLET_PREFIX_RE.match(line):
            return True
        return False

    def _split_candidate_tokens(self, text: str) -> list[str]:
        normalized = self._normalize_text(text)
        if not normalized:
            return []
        chunks = [normalized]
        for splitter in (PIPE_SPLIT_RE, COMMA_SPLIT_RE, SEMICOLON_SPLIT_RE):
            next_chunks: list[str] = []
            for chunk in chunks:
                next_chunks.extend(part for part in splitter.split(chunk) if part)
            chunks = next_chunks
        return [chunk.strip() for chunk in chunks if chunk.strip()]

    def _looks_like_contact_line(self, line: str) -> bool:
        return bool(EMAIL_RE.search(line) or PHONE_RE.search(line) or URL_RE.search(line))

    def _looks_like_role_line(self, line: str) -> bool:
        lowered = line.casefold()
        return any(keyword in lowered for keyword in ROLE_HINT_KEYWORDS)

    def _looks_like_company_line(self, lowered_line: str) -> bool:
        return any(keyword in lowered_line for keyword in {
            "inc",
            "llc",
            "ltd",
            "gmbh",
            "company",
            "solutions",
            "systems",
            "software",
            "technologies",
            "labs",
            "group",
            "holding",
        })

    def _looks_like_school_line(self, lowered_line: str) -> bool:
        return any(keyword in lowered_line for keyword in DEGREE_HINTS | {
            "school",
            "faculty",
            "lycee",
            "lycée",
        })

    def _looks_like_degree_line(self, lowered_line: str) -> bool:
        return any(keyword in lowered_line for keyword in DEGREE_HINTS)

    def _looks_like_field_of_study_line(self, lowered_line: str) -> bool:
        field_hints = {
            "engineering",
            "computer science",
            "electrical",
            "mechanical",
            "business",
            "economics",
            "mathematics",
            "physics",
            "chemistry",
            "psychology",
            "design",
        }
        return any(hint in lowered_line for hint in field_hints)

    def _looks_like_language_line(self, line: str) -> bool:
        lowered = line.casefold()
        return any(language in lowered for language in KNOWN_LANGUAGE_NAMES)

    def _looks_like_experience_line(self, line: str) -> bool:
        return bool(DATE_RANGE_RE.search(line) or self._looks_like_role_line(line))

    def _looks_like_education_line(self, line: str) -> bool:
        lowered = line.casefold()
        return self._looks_like_school_line(lowered) or self._looks_like_degree_line(lowered)

    def _looks_like_section_heading_candidate(self, line: str) -> bool:
        return self._detect_section_name(line) is not None

    def _derive_headline_from_summary(self, summary: str) -> str | None:
        first_sentence = re.split(r"[.!?]\s+", summary, maxsplit=1)[0].strip()
        if self._looks_like_role_line(first_sentence) and len(first_sentence) <= MAX_HEADLINE_LENGTH:
            return first_sentence
        return None

    def _compute_quality_score(
        self,
        *,
        text: str,
        page_count: int,
        used_ocr: bool,
        profile: CVStructuredProfile,
    ) -> int:
        text_length = len(text)
        line_count = self._line_count(text)
        alpha_chars = sum(1 for char in text if char.isalpha())
        digit_chars = sum(1 for char in text if char.isdigit())
        total_chars = max(1, len(text))
        alpha_ratio = alpha_chars / total_chars
        digit_ratio = digit_chars / total_chars

        score = 0
        if text_length >= 2000:
            score += 35
        elif text_length >= 900:
            score += 28
        elif text_length >= 300:
            score += 18
        elif text_length >= 120:
            score += 10
        else:
            score += 4

        if line_count >= 40:
            score += 10
        elif line_count >= 15:
            score += 6

        if page_count >= 2:
            score += 6
        elif page_count == 1:
            score += 3

        if 0.55 <= alpha_ratio <= 0.9:
            score += 12
        elif 0.4 <= alpha_ratio <= 0.95:
            score += 8
        else:
            score += 3

        if digit_ratio <= 0.25:
            score += 4

        if profile.personal_details.full_name:
            score += 4
        if profile.personal_details.headline:
            score += 4
        if profile.summary:
            score += 6
        if profile.skills:
            score += 6
        if profile.experience_entries:
            score += 8
        if profile.education_entries:
            score += 5
        if profile.language_entries:
            score += 3

        if used_ocr:
            score -= 12

        return max(0, min(score, 100))

    def _quality_label_from_score(self, quality_score: int) -> QualityLabel:
        if quality_score >= 70:
            return "high"
        if quality_score >= 40:
            return "medium"
        return "low"

    def _derive_status(
        self,
        *,
        normalized_text: str,
        warnings: Sequence[ParseWarning],
        profile: CVStructuredProfile,
    ) -> ParseStatus:
        if not normalized_text:
            return "empty"
        if any(item.severity == "error" for item in warnings):
            return "partial"
        meaningful_signal_count = sum(
            bool(value)
            for value in (
                profile.personal_details.full_name,
                profile.personal_details.headline,
                profile.summary,
                profile.skills,
                profile.experience_entries,
                profile.education_entries,
            )
        )
        if meaningful_signal_count >= 3:
            return "parsed"
        return "partial"

    def _normalize_text(self, text: str | None) -> str:
        if not text:
            return ""
        normalized = text.replace("\x00", "")
        normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
        normalized = WHITESPACE_RE.sub(" ", normalized)
        normalized = MULTIBREAK_RE.sub("\n\n", normalized)
        normalized = re.sub(r"\n +", "\n", normalized)
        normalized = re.sub(r" +\n", "\n", normalized)
        normalized = normalized.strip()
        return normalized

    def _extract_unique_matches(
        self,
        pattern: re.Pattern[str],
        lines: Iterable[str],
    ) -> list[str]:
        values: list[str] = []
        for line in lines:
            values.extend(match.group(0).strip() for match in pattern.finditer(line))
        return self._deduplicate_strings(values)

    def _normalize_phones(self, values: Iterable[str]) -> list[str]:
        normalized: list[str] = []
        for value in values:
            digits = re.sub(r"[^\d+]", "", value)
            if len(re.sub(r"\D", "", digits)) >= 8:
                normalized.append(digits)
        return self._deduplicate_strings(normalized)

    def _deduplicate_strings(self, values: Iterable[str]) -> list[str]:
        result: list[str] = []
        seen: set[str] = set()
        for value in values:
            normalized = value.strip()
            if not normalized:
                continue
            key = normalized.casefold()
            if key in seen:
                continue
            seen.add(key)
            result.append(normalized)
        return result

    def _display_skill(self, value: str) -> str:
        normalized = value.strip()
        canonical_map = {
            "node.js": "Node.js",
            "next.js": "Next.js",
            "github actions": "GitHub Actions",
            "power bi": "Power BI",
            "ci/cd": "CI/CD",
            "c#": "C#",
            "c++": "C++",
            ".net": ".NET",
            "aws": "AWS",
            "gcp": "GCP",
            "sql": "SQL",
        }
        return canonical_map.get(normalized.casefold(), normalized)

    def _truncate_text(self, text: str, max_chars: int) -> str:
        if len(text) <= max_chars:
            return text
        return text[: max_chars - 1].rstrip() + "…"

    def _detect_file_type(
        self,
        path: Path,
        mime_type: str | None,
    ) -> DocumentFileType:
        extension = path.suffix.lower()
        if extension == ".pdf":
            return "pdf"
        if extension == ".docx":
            return "docx"
        if extension == ".png":
            return "png"
        if extension == ".jpg":
            return "jpg"
        if extension == ".jpeg":
            return "jpeg"
        if mime_type:
            if mime_type == "application/pdf":
                return "pdf"
            if (
                mime_type
                == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ):
                return "docx"
            if mime_type.startswith("image/"):
                return "image"
        return "unknown"

    def _detect_mime_type(self, path: Path) -> str | None:
        detected, _ = mimetypes.guess_type(path.name)
        return detected

    def _suffix_from_content_type(self, content_type: str | None) -> str:
        if not content_type:
            return ""
        if content_type == "application/pdf":
            return ".pdf"
        if (
            content_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            return ".docx"
        if content_type == "image/png":
            return ".png"
        if content_type == "image/jpeg":
            return ".jpg"
        return ""

    def _line_count(self, text: str) -> int:
        if not text:
            return 0
        return len([line for line in text.split("\n") if line.strip()])

    def _sha256_for_file(self, path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            while True:
                chunk = handle.read(1024 * 1024)
                if not chunk:
                    break
                digest.update(chunk)
        return digest.hexdigest()

    def _elapsed_ms(self, started_at: float) -> int:
        return int((time.perf_counter() - started_at) * 1000)

    def _utc_now_iso(self) -> str:
        return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


def _parse_file_worker(file_path: str, config_dict: dict[str, Any]) -> CVParseResult:
    engine = UniversalCVParserEngine(ParserConfig.model_validate(config_dict))
    return engine.parse_file(file_path)


def configure_logging(level: int = logging.INFO) -> None:
    """Configure a production-friendly default logger for the parser engine."""
    logging.basicConfig(
        level=level,
        format="%(asctime)s %(levelname)s %(name)s %(message)s",
    )


def build_arg_parser() -> argparse.ArgumentParser:
    """Build a small CLI around the parser engine for local testing."""
    parser = argparse.ArgumentParser(
        description="Universal CV parser engine for PDF, DOCX, and image resumes.",
    )
    parser.add_argument("files", nargs="+", help="Paths to CV files.")
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Pretty-print JSON output.",
    )
    parser.add_argument(
        "--async-mode",
        action="store_true",
        help="Parse inputs concurrently using asyncio.",
    )
    parser.add_argument(
        "--multiprocess",
        action="store_true",
        help="Parse inputs with multiprocessing.",
    )
    parser.add_argument(
        "--ocr-languages",
        default="eng",
        help="OCR languages as a '+' separated list, e.g. 'eng+deu+tur'.",
    )
    return parser


async def _run_async_cli(
    engine: UniversalCVParserEngine,
    files: list[str],
    *,
    pretty: bool,
) -> None:
    results = await engine.parse_many_async(files)
    for result in results:
        print(result.model_dump_json(indent=2 if pretty else None))


def main(argv: Sequence[str] | None = None) -> int:
    """CLI entrypoint for manual testing and smoke validation."""
    configure_logging()
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    languages = tuple(
        token.strip() for token in args.ocr_languages.split("+") if token.strip()
    ) or DEFAULT_OCR_LANGUAGES
    engine = UniversalCVParserEngine(ParserConfig(ocr_languages=languages))

    if args.async_mode:
        asyncio.run(_run_async_cli(engine, list(args.files), pretty=args.pretty))
        return 0

    if args.multiprocess:
        results = engine.parse_many_multiprocess(list(args.files))
    else:
        results = [engine.parse_file(path) for path in args.files]

    for result in results:
        print(result.model_dump_json(indent=2 if args.pretty else None))
    return 0


if __name__ == "__main__":  # pragma: no cover - CLI entrypoint
    raise SystemExit(main())