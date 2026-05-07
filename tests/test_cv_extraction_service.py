from __future__ import annotations

from pathlib import Path

import pytest
from pypdf import PdfWriter

from hiring_radar.services import cv_extraction
from hiring_radar.services.cv_extraction import (
    CV_PARSE_STATUS_EMPTY,
    CV_PARSE_STATUS_PARSED,
    CvExtractionError,
    derive_cv_parse_status,
    extract_text_from_cv_file,
    normalize_extracted_text,
)

SIMPLE_TEXT_PDF_BYTES = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144]
/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length 55 >>
stream
BT
/F1 18 Tf
50 100 Td
(Hello Alice CV) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000010 00000 n 
0000000063 00000 n 
0000000122 00000 n 
0000000248 00000 n 
0000000353 00000 n 
trailer
<< /Root 1 0 R /Size 6 >>
startxref
423
%%EOF
"""


def test_normalize_extracted_text_collapses_whitespace() -> None:
    raw_text = "  Alice\r\n\r\n\r\nPython\t FastAPI  "
    assert normalize_extracted_text(raw_text) == "Alice\n\nPython FastAPI"
    assert normalize_extracted_text("   ") is None


def test_derive_cv_parse_status_distinguishes_text_and_empty() -> None:
    assert derive_cv_parse_status("Alice Example") == CV_PARSE_STATUS_PARSED
    assert derive_cv_parse_status(None) == CV_PARSE_STATUS_EMPTY


def test_extract_text_from_pdf_returns_parsed_status_for_text_pdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "alice_cv.pdf"
    pdf_path.write_bytes(SIMPLE_TEXT_PDF_BYTES)

    result = extract_text_from_cv_file(pdf_path)

    assert result.parse_status == CV_PARSE_STATUS_PARSED
    assert result.page_count == 1
    assert result.extracted_text == "Hello Alice CV"


def test_extract_text_from_pdf_returns_empty_status_for_blank_pdf(
    tmp_path: Path,
) -> None:
    pdf_path = tmp_path / "blank_cv.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=144)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    result = extract_text_from_cv_file(pdf_path)

    assert result.parse_status == CV_PARSE_STATUS_EMPTY
    assert result.page_count == 1
    assert result.extracted_text is None


def test_extract_text_from_docx_returns_parsed_status(tmp_path: Path) -> None:
    from docx import Document

    docx_path = tmp_path / "alice_cv.docx"
    document = Document()
    document.add_paragraph("Alice Example")
    document.add_paragraph("Senior Backend Engineer")
    document.add_paragraph("Python, FastAPI, SQL")
    document.save(docx_path)

    result = extract_text_from_cv_file(docx_path)

    assert result.parse_status == CV_PARSE_STATUS_PARSED
    assert result.page_count == 1
    assert result.extracted_text is not None
    assert "Alice Example" in result.extracted_text
    assert "Senior Backend Engineer" in result.extracted_text
    assert "Python, FastAPI, SQL" in result.extracted_text


def test_extract_text_from_png_returns_parsed_status(tmp_path: Path, monkeypatch) -> None:
    from PIL import Image

    png_path = tmp_path / "alice_cv.png"
    Image.new("RGB", (240, 120), color="white").save(png_path)

    monkeypatch.setattr(
        cv_extraction,
        "_ocr_image_to_text",
        lambda _image, *, lang: "Alice Example\nSenior Backend Engineer\nPython",
    )

    result = extract_text_from_cv_file(png_path)

    assert result.parse_status == CV_PARSE_STATUS_PARSED
    assert result.page_count == 1
    assert result.extracted_text == "Alice Example\nSenior Backend Engineer\nPython"


def test_extract_text_from_cv_file_rejects_malformed_legacy_doc_files(tmp_path: Path) -> None:
    doc_path = tmp_path / "alice_cv.doc"
    doc_path.write_bytes(b"fake-doc")

    with pytest.raises(CvExtractionError, match="does not look like a Word"):
        extract_text_from_cv_file(doc_path)


def test_normalize_extracted_text_uses_foundation_cleanup_rules() -> None:
    raw_text = "\ufeffAlice\u200b Example\r\n\r\n•••\r\nPython\t FastAPI\r\nPage 1 of 1"
    assert normalize_extracted_text(raw_text) == "Alice Example\n\nPython FastAPI"



def test_extract_text_from_pdf_prefers_layout_aware_reading_order(
    tmp_path: Path,
    monkeypatch,
) -> None:
    pdf_path = tmp_path / "alice_cv.pdf"
    pdf_path.write_bytes(SIMPLE_TEXT_PDF_BYTES)

    ordered_text = "Alice Example\n\nSenior Backend Engineer\n\nPython"
    monkeypatch.setattr(
        cv_extraction,
        "_extract_pdf_text_with_layout",
        lambda _path: cv_extraction.CvExtractionResult(
            extracted_text=ordered_text,
            parse_status=CV_PARSE_STATUS_PARSED,
            page_count=1,
        ),
    )

    result = extract_text_from_cv_file(pdf_path)

    assert result.extracted_text == ordered_text
    assert result.parse_status == CV_PARSE_STATUS_PARSED


def test_extract_text_from_pdf_falls_back_to_pypdf_when_layout_is_unavailable(
    tmp_path: Path,
    monkeypatch,
) -> None:
    pdf_path = tmp_path / "alice_cv.pdf"
    pdf_path.write_bytes(SIMPLE_TEXT_PDF_BYTES)

    monkeypatch.setattr(cv_extraction, "_extract_pdf_text_with_layout", lambda _path: None)

    result = extract_text_from_cv_file(pdf_path)

    assert result.extracted_text == "Hello Alice CV"
    assert result.page_count == 1


def test_extract_text_from_pdf_falls_back_to_ocr_when_text_is_thin(
    tmp_path: Path,
    monkeypatch,
) -> None:
    pdf_path = tmp_path / "scanned_cv.pdf"
    pdf_path.write_bytes(SIMPLE_TEXT_PDF_BYTES)

    ocr_text = (
        "Alice Example\n"
        "Senior Backend Engineer\n"
        "alice@example.com | Berlin | Remote\n\n"
        "Professional Summary\n"
        "Backend engineer specialised in FastAPI and data-intensive services.\n\n"
        "Technical Skills\n"
        "Python, FastAPI, SQL, Docker, AWS"
    )
    monkeypatch.setattr(cv_extraction, "_extract_pdf_text_with_layout", lambda _path: None)
    monkeypatch.setattr(
        cv_extraction,
        "_extract_pdf_text_with_ocr",
        lambda _path: cv_extraction.CvExtractionResult(
            extracted_text=ocr_text,
            parse_status=cv_extraction.CV_PARSE_STATUS_PARSED,
            page_count=1,
        ),
    )

    result = extract_text_from_cv_file(pdf_path)

    assert result.extracted_text == ocr_text
    assert result.parse_status == cv_extraction.CV_PARSE_STATUS_PARSED
    assert result.page_count == 1


def test_extract_text_from_pdf_gracefully_returns_pypdf_when_ocr_unavailable(
    tmp_path: Path,
    monkeypatch,
) -> None:
    pdf_path = tmp_path / "alice_cv.pdf"
    pdf_path.write_bytes(SIMPLE_TEXT_PDF_BYTES)

    monkeypatch.setattr(cv_extraction, "_extract_pdf_text_with_layout", lambda _path: None)
    monkeypatch.setattr(cv_extraction, "_extract_pdf_text_with_ocr", lambda _path: None)

    result = extract_text_from_cv_file(pdf_path)

    assert result.extracted_text == "Hello Alice CV"
    assert result.page_count == 1


def test_extract_text_from_docx_includes_section_headers_and_tables(
    tmp_path: Path,
) -> None:
    from docx import Document

    docx_path = tmp_path / "alice_cv.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Alice Example"
    document.add_paragraph("Senior Backend Engineer")
    document.add_paragraph("Technical Skills")
    skill_table = document.add_table(rows=1, cols=2)
    skill_table.rows[0].cells[0].text = "Backend"
    skill_table.rows[0].cells[1].text = "Python, FastAPI"
    document.sections[0].footer.paragraphs[0].text = "alice@example.com"
    document.save(docx_path)

    result = extract_text_from_cv_file(docx_path)

    assert result.parse_status == CV_PARSE_STATUS_PARSED
    assert result.extracted_text is not None
    assert "Alice Example" in result.extracted_text
    assert "Senior Backend Engineer" in result.extracted_text
    assert "Backend | Python, FastAPI" in result.extracted_text
    assert "alice@example.com" in result.extracted_text


def test_extract_text_from_image_tries_multiple_psm_modes(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from PIL import Image

    png_path = tmp_path / "alice_cv.png"
    Image.new("RGB", (400, 200), color="white").save(png_path)

    calls: list[str] = []

    def fake_image_to_string(_image, *, lang, config="", **_kwargs):
        calls.append(config)
        if "--psm 6" in config:
            return "Alice Example\nPython"
        if "--psm 4" in config:
            return "Alice Example\nSenior Backend Engineer\nPython, FastAPI, SQL"
        return "noise"

    class _FakeTesseractNotFound(Exception):
        pass

    fake_tesseract = type(
        "FakeTesseract",
        (),
        {"image_to_string": staticmethod(fake_image_to_string)},
    )

    import sys
    monkeypatch.setitem(sys.modules, "pytesseract", fake_tesseract)
    monkeypatch.setitem(
        sys.modules,
        "pytesseract.pytesseract",
        type(
            "_FakeErrors",
            (),
            {"TesseractNotFoundError": _FakeTesseractNotFound},
        ),
    )

    # Patch orientation detection + language detection to stay deterministic.
    monkeypatch.setattr(cv_extraction, "_correct_image_orientation", lambda image: image)
    monkeypatch.setattr(cv_extraction, "_detect_ocr_languages", lambda: "eng")

    # pytesseract attribute for TesseractNotFoundError import inside function
    monkeypatch.setattr(
        cv_extraction,
        "_ocr_image_to_text",
        cv_extraction._ocr_image_to_text,
    )

    # Inject the submodule the function imports from: `from pytesseract import TesseractNotFoundError`
    fake_tesseract.TesseractNotFoundError = _FakeTesseractNotFound

    result = extract_text_from_cv_file(png_path)

    assert result.parse_status == CV_PARSE_STATUS_PARSED
    assert result.extracted_text is not None
    assert "Senior Backend Engineer" in result.extracted_text
    assert any("--psm 4" in config for config in calls)


def test_column_split_detects_right_sidebar_layout() -> None:
    words = []
    for index, top in enumerate(range(100, 700, 40)):
        words.append({"text": f"left{index}", "left": 80, "top": top, "width": 80, "height": 20, "cx": 120})
        words.append({"text": f"leftb{index}", "left": 720, "top": top, "width": 100, "height": 20, "cx": 770})
    for index, top in enumerate(range(150, 600, 50)):
        words.append({"text": f"right{index}", "left": 1740, "top": top, "width": 100, "height": 20, "cx": 1790})

    split_x, is_two_column = cv_extraction._detect_column_split(words, 2400)

    assert is_two_column is True
    assert 1200 < split_x < 1740


def test_extract_text_from_webp_returns_parsed_status(tmp_path: Path, monkeypatch) -> None:
    from PIL import Image

    webp_path = tmp_path / "leyla_cv.webp"
    Image.new("RGB", (240, 120), color="white").save(webp_path, "WEBP")

    monkeypatch.setattr(
        cv_extraction,
        "_ocr_image_to_text",
        lambda _image, *, lang: "Leyla Aydın\nProje Mühendisi\nmerhaba@example.com",
    )

    result = extract_text_from_cv_file(webp_path)

    assert result.parse_status == CV_PARSE_STATUS_PARSED
    assert result.page_count == 1
    assert result.extraction_method == cv_extraction.CV_EXTRACTION_METHOD_IMAGE_OCR
    assert result.used_ocr is True
    assert "Leyla Aydın" in (result.extracted_text or "")


def test_late_two_column_reconstruction_keeps_profile_full_width() -> None:
    words = [
        {"text": "Leyla", "left": 80, "top": 60, "width": 60, "height": 20, "cx": 110},
        {"text": "Aydın", "left": 150, "top": 60, "width": 70, "height": 20, "cx": 185},
        {"text": "Profil", "left": 80, "top": 140, "width": 60, "height": 20, "cx": 110},
        {"text": "karmaşık", "left": 80, "top": 180, "width": 85, "height": 20, "cx": 122},
        {"text": "projelerini", "left": 175, "top": 180, "width": 95, "height": 20, "cx": 222},
        {"text": "tamamlamaya", "left": 280, "top": 180, "width": 110, "height": 20, "cx": 335},
        {"text": "kadar", "left": 400, "top": 180, "width": 60, "height": 20, "cx": 430},
        {"text": "İŞ", "left": 80, "top": 260, "width": 35, "height": 20, "cx": 98},
        {"text": "DENEYİMİ", "left": 125, "top": 260, "width": 100, "height": 20, "cx": 175},
        {"text": "EĞİTİM", "left": 650, "top": 260, "width": 90, "height": 20, "cx": 695},
    ]
    for idx, top in enumerate(range(310, 630, 40)):
        words.append({"text": f"exp{idx}", "left": 80, "top": top, "width": 60, "height": 20, "cx": 110})
        words.append({"text": f"edu{idx}", "left": 650, "top": top, "width": 60, "height": 20, "cx": 680})

    text = cv_extraction._reconstruct_late_two_column_text(words, 900)

    assert text is not None
    assert "karmaşık projelerini tamamlamaya kadar" in text
    assert text.index("İŞ DENEYİMİ") < text.index("exp0") < text.index("EĞİTİM")
    assert text.index("EĞİTİM") < text.index("edu0")
